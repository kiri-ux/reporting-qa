"""What the reporting tracker's market codes stand for.

"LOCK KNOX" is Lockwood Digital Solutions Knoxville. "INNO" is Innovision
Advertising. Those codes reach the board when a row is approved for a client
the order export has never heard of: there is no order line to take a partner
from, so the tracker's own market column is all there is, and unresolved it
became a partner card of its own beside the real one.

The list is columns H and J of Partner Onboarding.xlsx - the abbreviation and
the market name - which is where the codes are decided. Guessing from the
letters got LOCK KNOX right and would not have got 3P or 270M right at all.

An abbreviation that stands for two different markets is left out. There are
three, and half an answer is worse than none for this.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

SEED = Path(__file__).resolve().parent / "seed" / "market_abbr.json"


def key(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "")).strip().upper()


def _load() -> dict[str, str]:
    # A LATER UPLOAD BEATS THE SHIPPED COPY. Markets are onboarded every month
    # and the file that decides this lives in somebody's Drive, not in here.
    from .config import settings

    out: dict[str, str] = {}
    for src in (SEED, settings.data_dir / "market_abbr.json"):
        try:
            data = json.loads(Path(src).read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        if isinstance(data, dict):
            out.update({key(k): str(v).strip() for k, v in data.items()
                        if key(k) and str(v).strip()})
    return out


_MAP: dict[str, str] | None = None


def market_for(code: str) -> str:
    """The market this code names, or "" - never a guess."""
    global _MAP
    if _MAP is None:
        _MAP = _load()
    return _MAP.get(key(code), "")


def reload() -> int:
    """Re-read after an upload. Returns how many codes are known."""
    global _MAP
    _MAP = _load()
    return len(_MAP)


def read_sheet(raw: bytes) -> dict[str, str]:
    """Columns H and J of the onboarding workbook: abbreviation, market name.

    Read by position, not by header, because the header row of that sheet is
    a design that changes and the columns are the thing somebody points at.
    An abbreviation appearing against two different markets is dropped.
    """
    import io

    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    seen: dict[str, set[str]] = {}
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=10, values_only=True):
        abbr = key(str(row[7] or ""))
        name = re.sub(r"\s+", " ", str(row[9] or "")).strip()
        if abbr and name:
            seen.setdefault(abbr, set()).add(name)
    return {a: next(iter(v)) for a, v in seen.items() if len(v) == 1}
